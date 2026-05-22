"""Deep validation of NRO2026.xlsx."""
import json
import sys
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

PATH = Path(r"C:\Users\ASUS\Desktop\NRO2026.xlsx")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from app.panne_type_catalog import BERCEAU_PANNE_TYPES  # noqa: E402

OUT = Path(__file__).resolve().parent.parent / "nro2026_deep_report.json"

IMAGE_LIST_53 = [
    "Amorçage", "Arrêt", "Arrêt de convoyeur", "Arrêt programmé",
    "Basculement de A1 à A3", "Basculement de A3 à A1",
    "Blocage de coulisseau", "Blocage de fil", "Blocage de pilote",
    "Blocage de poinçons", "Blocage de robot de transfert", "Blocage de serrage",
    "Blocage manipulateur", "Cable de masse", "Changement de bobine de fil",
    "Changement de diffuseur", "Changement de torche", "Changement des éléectrodes",
    "Chute de la torche", "Chute de pièce", "Collision", "Décyclage",
    "Défaut de robot MAG", "Dégradation de gaine", "Fil collé", "Fuite d'air",
    "Fuite d'eau", "Lancement de soudage", "Ligne hors service", "Ligne vide",
    "Manque de débit d'eau", "Manque de référence", "Manque d'emballage", "MES",
    "Nettoyage des buses", "Perte d'énergie", "Poinçonnage", "Probleme automate",
    "Problème dans circuit de sécurité", "Problème d'ascenseur", "Problème de bridage",
    "Problème de capteur", "Problème de générateur", "Problème de lubrification",
    "Problème de soudage", "Problème de vérin", "Problème élévateur", "Problème GAZ",
    "Problème maquette", "Problème pression", "Problème teach", "Réglage Z-Citrou",
    "Remplissage des réservoirs de lubrification",
]


def norm(s):
    return (s or "").strip()


def parse_date(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    s = norm(str(v))
    if " " in s:
        s = s.split(" ")[0]
    return s


def main():
    wb = load_workbook(PATH, read_only=True, data_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header = [norm(c) for c in rows[0]]
    idx = {h: i for i, h in enumerate(header)}
    records = []
    for row in rows[1:]:
        if not any(row):
            continue
        rec = {h: row[idx[h]] if idx[h] < len(row) else None for h in header}
        records.append(rec)

    problems = [norm(r.get("Problème")) for r in records]
    problem_counts = Counter(problems)
    excel_problems = set(problems)

    catalog = set(BERCEAU_PANNE_TYPES)
    image = set(IMAGE_LIST_53)

    # Typo / mapping notes
    TYPO_MAP = {
        "Changement des éléectrodes": "Changement des électrodes",
        "Problème élevateur": "Problème élévateur",
        "Probleme de Arrêt": "Arrêt (typo Excel)",
        "Probleme GAZ": "Problème GAZ",
        "Probleme de capteur": "Problème de capteur",
        "Probleme de vérin": "Problème de vérin",
        "Probleme de générateur": "Problème de générateur",
    }

    in_excel_not_catalog = sorted(excel_problems - catalog)
    in_catalog_not_excel = sorted(catalog - excel_problems)
    in_image_not_excel = sorted(image - excel_problems)
    in_excel_not_image = sorted(excel_problems - image)

    empty_moyen = [r for r in records if not norm(str(r.get("Moyen") or ""))]
    invalid_temps = []
    for i, r in enumerate(records, start=2):
        t = r.get("Temps (min)")
        try:
            v = float(t) if t is not None else None
            if v is None or v < 0:
                invalid_temps.append((i, t))
        except (TypeError, ValueError):
            invalid_temps.append((i, t))

    dates = sorted(set(parse_date(r["Date"]) for r in records))
    by_date = Counter(parse_date(r["Date"]) for r in records)
    by_type = Counter(norm(r.get("Type")) for r in records)
    by_module = Counter(norm(r.get("Module")) for r in records)
    by_diversite = Counter(norm(r.get("Diversité")) for r in records)
    total_min = sum(int(float(r["Temps (min)"])) for r in records)

    # module-poste-moyen combos
    mpm = Counter(
        (norm(r["Module"]), norm(r["Poste"]), norm(str(r.get("Moyen") or "")))
        for r in records
    )

    report = {
        "summary": {
            "total_records": len(records),
            "total_downtime_minutes": total_min,
            "date_range": {"first": dates[0], "last": dates[-1], "distinct_days": len(dates)},
            "missing_calendar_days_may_2026": [
                d for d in range(1, 17)
                if f"2026-05-{d:02d}" not in dates
            ],
            "records_per_day": dict(sorted(by_date.items())),
        },
        "columns": header,
        "distributions": {
            "Type": dict(by_type.most_common()),
            "Module": dict(by_module.most_common()),
            "Diversité": dict(by_diversite.most_common()),
        },
        "probleme": {
            "unique_in_excel": len(excel_problems),
            "all_with_counts": dict(problem_counts.most_common()),
            "in_excel_not_in_app_catalog": in_excel_not_catalog,
            "in_app_catalog_not_used_in_excel": in_catalog_not_excel,
            "in_reference_image_not_in_excel": in_image_not_excel,
            "in_excel_not_in_reference_image": in_excel_not_image,
            "known_typos_in_excel": {k: v for k, v in TYPO_MAP.items() if k in excel_problems},
        },
        "moyen": {
            "empty_count": len(empty_moyen),
            "empty_pct": round(100 * len(empty_moyen) / len(records), 1),
            "empty_by_module_poste": {
                f"{m} | {p}": n
                for (m, p), n in Counter(
                    (norm(r["Module"]), norm(r["Poste"])) for r in empty_moyen
                ).items()
            },
        },
        "matrix_combos_top30": [
            {"module": k[0], "poste": k[1], "moyen": k[2] or "(vide)", "count": v}
            for k, v in mpm.most_common(30)
        ],
        "data_quality": {
            "invalid_temps_rows": invalid_temps[:20],
            "invalid_temps_count": len(invalid_temps),
        },
    }

    OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
