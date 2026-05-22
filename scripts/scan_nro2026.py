"""Scan NRO2026.xlsx workbook — full structure report."""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

PATH = Path(r"C:\Users\ASUS\Desktop\NRO2026.xlsx")
OUT = Path(__file__).resolve().parent.parent / "nro2026_scan_report.json"


def cell_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def scan_sheet(ws, max_rows=10000):
    all_rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= max_rows:
            break
        all_rows.append(row)

    max_cols = max((len(r) for r in all_rows if r), default=0)
    padded = []
    for r in all_rows:
        cells = list(r) if r else []
        while len(cells) < max_cols:
            cells.append(None)
        padded.append(tuple(cell_str(cells[j]) for j in range(max_cols)))

    header = None
    data_start = 0
    for idx, row in enumerate(padded[:25]):
        non_empty = sum(1 for c in row if c)
        if non_empty >= 2:
            header = [c if c else f"col_{j}" for j, c in enumerate(row)]
            data_start = idx + 1
            break

    data_rows = [r for r in padded[data_start:] if any(c for c in r)]

    info = {
        "rows_read": len(all_rows),
        "data_rows_non_empty": len(data_rows),
        "max_columns": max_cols,
        "header_row_index": data_start - 1 if header else None,
        "headers": header,
        "first_8_data_rows": [list(r[: min(20, max_cols)]) for r in data_rows[:8]],
        "last_5_data_rows": [list(r[: min(20, max_cols)]) for r in data_rows[-5:]],
    }

    if header and data_rows:
        col_stats = {}
        for j, h in enumerate(header):
            if h.startswith("col_") and not any(
                cell_str(data_rows[k][j]) for k in range(min(5, len(data_rows))) if j < len(data_rows[k])
            ):
                continue
            vals = []
            for r in data_rows:
                if j < len(r) and r[j]:
                    vals.append(r[j])
            unique = sorted(set(vals), key=lambda x: (len(x), x.lower()))
            col_stats[h] = {
                "non_empty_count": len(vals),
                "unique_count": len(unique),
                "unique_values": unique
                if len(unique) <= 120
                else unique[:120] + [f"... +{len(unique) - 120} more"],
            }
        info["column_stats"] = col_stats

    return info


def main():
    if not PATH.exists():
        print(f"File not found: {PATH}", file=sys.stderr)
        sys.exit(1)

    wb = load_workbook(PATH, read_only=True, data_only=True)
    report = {
        "file": str(PATH),
        "size_bytes": PATH.stat().st_size,
        "sheet_count": len(wb.sheetnames),
        "sheet_names": wb.sheetnames,
        "sheets": {},
    }

    for name in wb.sheetnames:
        report["sheets"][name] = scan_sheet(wb[name])

    wb.close()
    OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Wrote {OUT}")
    print(f"Sheets: {report['sheet_names']}")
    for sn, si in report["sheets"].items():
        print(f"  [{sn}] rows={si['data_rows_non_empty']} cols={si['max_columns']}")
        if si.get("headers"):
            print(f"    headers: {si['headers'][:12]}{'...' if len(si['headers'])>12 else ''}")


if __name__ == "__main__":
    main()
