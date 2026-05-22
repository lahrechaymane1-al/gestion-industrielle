"""Audit NRO2026 imports vs Excel and dashboard prerequisites."""
from __future__ import annotations

import os
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "core.settings")

import django

django.setup()

from django.db.models import Sum
from openpyxl import load_workbook

from app.downtime_impact import (
    berceau_alerte_impact_pct,
    parse_diversite_from_cause,
    production_rows_map_for_berceau_alertes,
    total_objectif_all_hours,
)
from arret.models import AlertePanne
from production.models import ProductionBerceau

XLSX = Path(r"C:\Users\ASUS\Desktop\NRO2026.xlsx")


def load_excel_by_day() -> dict[date, dict]:
    wb = load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    hdr = [str(c or "").strip() for c in rows[0]]
    col = {h: i for i, h in enumerate(hdr)}

    def cell(r, name):
        i = col[name]
        return r[i] if i < len(r) else None

    out: dict[date, dict] = defaultdict(lambda: {"lines": 0, "min": 0})
    for r in rows[1:]:
        if not any(r):
            continue
        d = cell(r, "Date")
        if hasattr(d, "date"):
            d = d.date()
        else:
            d = date.fromisoformat(str(d).split()[0])
        m = int(float(cell(r, "Temps (min)") or 0))
        out[d]["lines"] += 1
        out[d]["min"] += m
    return dict(out)


def main() -> int:
    excel_by_day = load_excel_by_day()
    issues: list[str] = []

    print("=== IMPORT vs EXCEL (minutes must match) ===")
    for d in sorted(excel_by_day):
        ex = excel_by_day[d]
        tag = f"[import:NRO2026]:{d.isoformat()}"
        qs = AlertePanne.objects.filter(
            equipe="Berceau", date=d, cause__startswith=tag, is_deleted=False
        )
        db_n = qs.count()
        db_min = int(qs.aggregate(s=Sum("temps_arret_min"))["s"] or 0)
        ok_min = db_min == ex["min"]
        print(
            f"  {d}  excel {ex['lines']} lignes / {ex['min']} min  ->  "
            f"db {db_n} fiches / {db_min} min  [{'OK' if ok_min else 'ERREUR'}]"
        )
        if not ok_min:
            issues.append(f"{d}: excel {ex['min']} min vs db {db_min} min")

    d16 = date(2026, 5, 16)
    ex16 = excel_by_day.get(d16)
    if ex16:
        qs16 = AlertePanne.objects.filter(
            equipe="Berceau",
            date=d16,
            cause__startswith="[import:NRO2026]:2026-05-16",
            is_deleted=False,
        )
        m16 = int(qs16.aggregate(s=Sum("temps_arret_min"))["s"] or 0)
        ok = m16 == ex16["min"]
        print(
            f"  2026-05-16 (Excel 16/05)  excel {ex16['lines']}L/{ex16['min']}m  ->  "
            f"db {qs16.count()}F/{m16}m  [{'OK' if ok else 'ERREUR'}]"
        )
        if not ok:
            issues.append(f"2026-05-16: excel16 {ex16['min']} vs db {m16}")

    print("\n=== PRODUCTION vs ARRÊTS (dashboard) ===")
    check_days = sorted(set(excel_by_day))
    for d in check_days:
        al_shifts = set(
            AlertePanne.objects.filter(equipe="Berceau", date=d, is_deleted=False).values_list(
                "shift", flat=True
            )
        )
        prod_rows = list(ProductionBerceau.objects.filter(date=d))
        pr_shifts = [p.shift for p in prod_rows]
        obj = total_objectif_all_hours(prod_rows)
        warn = ""
        if al_shifts and pr_shifts and not (al_shifts & set(pr_shifts)):
            warn = "  <- arrêts et prod pas même shift (dashboard utilise repli prod jour)"
        print(f"  {d}  arrêts {sorted(al_shifts)}  prod {pr_shifts}  objectif_jour={obj}{warn}")

    print("\n=== PARETO 2026-05-16 (arrêts shift A) ===")
    alertes = list(
        AlertePanne.objects.filter(date=d16, equipe="Berceau", shift="A", is_deleted=False)
    )
    prod_map = production_rows_map_for_berceau_alertes(alertes)
    prod_a = prod_map.get((d16, "A"), [])
    prod_all = list(ProductionBerceau.objects.filter(date=d16))
    prod_use = prod_all if prod_all else prod_a
    total_pct = sum(berceau_alerte_impact_pct(a, prod_use) for a in alertes)
    total_min = sum(a.temps_arret_min for a in alertes)
    no_tag = sum(1 for a in alertes if parse_diversite_from_cause(a.cause) is None)
    zero_min = sum(1 for a in alertes if a.temps_arret_min <= 0)
    print(f"  fiches: {len(alertes)}, minutes: {total_min}, sum impact%: {round(total_pct, 2)}")
    print(f"  prod shift A: {len(prod_a)}  prod jour: {len(prod_all)}  objectif utilisé: {total_objectif_all_hours(prod_use)}")
    print(f"  sans tag diversité: {no_tag}  fiches 0 min: {zero_min}")

    print("\n=== HEURES PRODUCTION (1-8) ===")
    bad_h = AlertePanne.objects.filter(
        equipe="Berceau", cause__startswith="[import:NRO2026]:", is_deleted=False
    ).exclude(heure_production__gte=1, heure_production__lte=8)
    print(f"  fiches hors H1-H8: {bad_h.count()}")

    print("\n=== RÉSUMÉ ===")
    if issues:
        print("  PROBLÈMES:")
        for i in issues:
            print(f"    - {i}")
        return 1
    print("  Minutes Excel = base OK pour tous les jours importés.")
    print("  Dashboard 16/05: OK si objectif_jour > 0 (prod A+B+N du jour).")
    return 0


if __name__ == "__main__":
    sys.exit(main())
